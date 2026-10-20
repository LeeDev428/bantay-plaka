from datetime import datetime, timezone as datetime_timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.test import RequestFactory, SimpleTestCase
from openpyxl import load_workbook

from apps.export_helpers import build_excel_response, build_export_chart_data, build_pdf_response
from apps.logs.export_views import export_logs_excel, export_logs_pdf


class VehicleLogExportChartTests(SimpleTestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.user = SimpleNamespace(
            is_authenticated=True,
            is_resident=lambda: False,
        )
        self.logs = [
            SimpleNamespace(
                status='TIME_IN',
                timestamp=datetime(2026, 10, 10, 3, 0, tzinfo=datetime_timezone.utc),
                plate_number='ABC1234',
                entry_type='VISITOR',
                source='CAMERA',
                camera_role='ENTRY_CAM',
                resident_name='',
                visitor_name='Guest',
                logged_by=None,
                get_entry_type_display=lambda: 'Visitor',
                get_source_display=lambda: 'Camera',
                get_camera_role_display=lambda: 'Entry Camera',
            ),
            SimpleNamespace(
                status='TIME_OUT',
                timestamp=datetime(2026, 10, 10, 4, 0, tzinfo=datetime_timezone.utc),
                plate_number='XYZ9876',
                entry_type='RESIDENT',
                source='CAMERA',
                camera_role='EXIT_CAM',
                resident_name='Resident',
                visitor_name='',
                logged_by=None,
                get_entry_type_display=lambda: 'Resident',
                get_source_display=lambda: 'Camera',
                get_camera_role_display=lambda: 'Exit Camera',
            ),
            SimpleNamespace(
                status='TIME_IN',
                timestamp=datetime(2026, 10, 11, 3, 0, tzinfo=datetime_timezone.utc),
                plate_number='ABC1234',
                entry_type='VISITOR',
                source='CAMERA',
                camera_role='ENTRY_CAM',
                resident_name='',
                visitor_name='Guest',
                logged_by=None,
                get_entry_type_display=lambda: 'Visitor',
                get_source_display=lambda: 'Camera',
                get_camera_role_display=lambda: 'Entry Camera',
            ),
        ]
        self.queryset_patch = patch(
            'apps.logs.export_views.VehicleLog.objects.select_related',
            return_value=SimpleNamespace(order_by=lambda *args: self.logs),
        )
        self.queryset_patch.start()
        self.addCleanup(self.queryset_patch.stop)
        self.blacklist_patch = patch(
            'apps.logs.export_views.get_active_blacklist_map',
            return_value={},
        )
        self.blacklist_patch.start()
        self.addCleanup(self.blacklist_patch.stop)

    def _request(self, path):
        request = self.factory.get(path)
        request.user = self.user
        return request

    def test_shared_chart_data_counts_categories_and_days(self):
        chart_data = build_export_chart_data(
            [
                ['Admin', '2026-10-10'],
                ['Guard', '2026-10-10'],
                ['Guard', '2026-10-11'],
                ['—', '—'],
            ],
            0,
            1,
            'Users by Role',
        )

        self.assertEqual(chart_data['category'], [('Admin', 1), ('Guard', 2)])
        self.assertEqual(chart_data['category_title'], 'Users by Role')
        self.assertEqual(chart_data['daily'], [('2026-10-10', 2), ('2026-10-11', 1)])

    def test_excel_export_includes_both_charts(self):
        response = export_logs_excel(self._request('/logs/export/excel/'))
        workbook = load_workbook(BytesIO(response.content))
        summary = workbook['Summary']

        self.assertEqual(len(summary._charts), 2)
        self.assertEqual(summary['A26'].value, 'Time In')
        self.assertEqual(summary['B26'].value, 2)
        self.assertEqual(summary['A27'].value, 'Time Out')
        self.assertEqual(summary['B27'].value, 1)
        self.assertEqual(summary['D26'].value, '2026-10-10')
        self.assertEqual(summary['E26'].value, 2)

    def test_excel_export_shows_na_when_no_records_are_available(self):
        self.logs.clear()

        response = export_logs_excel(self._request('/logs/export/excel/'))
        workbook = load_workbook(BytesIO(response.content))
        summary = workbook['Summary']

        self.assertEqual(summary._charts, [])
        self.assertEqual(summary['D2'].value, 'Time In vs. Time Out: N/A')
        self.assertEqual(summary['L2'].value, 'Records per Day: N/A')

    def test_pdf_export_builds_with_both_charts(self):
        response = export_logs_pdf(self._request('/logs/export/pdf/'))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertGreater(len(response.content), 1000)

    def test_pdf_export_builds_na_visuals_when_no_records_are_available(self):
        self.logs.clear()

        response = export_logs_pdf(self._request('/logs/export/pdf/'))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertGreater(len(response.content), 1000)


class ExportTableSizingTests(SimpleTestCase):
    def test_excel_columns_expand_for_long_headers_and_values(self):
        response = build_excel_response(
            'sizing_test',
            'Sizing Test',
            headers=['Record identifier', 'Description'],
            col_widths=[5, 8],
            rows=[[1, 'A description that should widen its column']],
        )
        worksheet = load_workbook(BytesIO(response.content)).active

        self.assertGreaterEqual(worksheet.column_dimensions['A'].width, len('Record identifier') + 2)
        self.assertGreaterEqual(
            worksheet.column_dimensions['B'].width,
            len('A description that should widen its column') + 2,
        )
        self.assertTrue(worksheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(worksheet.page_setup.fitToWidth, 1)
        self.assertEqual(worksheet.page_setup.orientation, 'landscape')

    def test_pdf_table_fits_landscape_page_and_wraps_cell_content(self):
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm

        document_instance = Mock()
        document_instance.leftMargin = 12 * mm
        document_instance.rightMargin = 12 * mm
        with patch('reportlab.platypus.SimpleDocTemplate', return_value=document_instance):
            response = build_pdf_response(
                'sizing_test',
                'Sizing Test',
                headers=['Name', 'Description'],
                col_widths_mm=[10, 10],
                rows=[['Resident', 'Long text with <markup> that should wrap within the page']],
                mono_col_idx=0,
            )

        self.assertEqual(response.status_code, 200)
        table = document_instance.build.call_args.args[0][-1]
        available_width = landscape(A4)[0] - 24 * mm
        self.assertAlmostEqual(sum(table._argW), available_width)
        self.assertEqual(table._cellvalues[1][0].style.fontName, 'Courier-Bold')
        self.assertEqual(table._cellvalues[1][1].text, 'Long text with &lt;markup&gt; that should wrap within the page')
